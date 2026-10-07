import asyncio
import logging
import os
import re
import sys
from pathlib import Path
from dataclasses import dataclass
from typing import List, Optional, Set, Tuple
from urllib.parse import urljoin

# Loyiha ildiz papkasini sys.path ga qo'shish
root_dir = Path(__file__).resolve().parent.parent
if str(root_dir) not in sys.path:
    sys.path.insert(0, str(root_dir))

import httpx
from bs4 import BeautifulSoup
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from controlpribor_bot.database.db_manager import DatabaseManager

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
logger = logging.getLogger(__name__)

# Top 16 mashhur brendlar lug'ati
KNOWN_BRANDS = [
    "Fluke",
    "Chint",
    "Akfix",
    "Mastech",
    "UNI-T",
    "Testo",
    "RGK",
    "Victor",
    "Trotec",
    "Michell",
    "АКИП",
    "ОВЕН",
    "CEM",
    "Мегеон",
    "Sonel",
    "APPA",
]


@dataclass
class ScrapedProduct:
    name: str
    price: str
    category: str
    brand: str
    image_url: Optional[str]
    product_url: str


class ControlPriborScraper:
    def __init__(self, base_url: str = "https://controlpribor.uz", max_concurrency: int = 5):
        self.base_url = base_url.rstrip("/")
        self.categories_url = f"{self.base_url}/categories/"
        self.max_concurrency = max_concurrency
        self.headers = {
            "User-Agent": (
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                "(KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
            ),
            "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
            "Accept-Language": "ru-RU,ru;q=0.9,uz;q=0.8,en;q=0.7",
        }

    def extract_brand_name(self, title: str) -> str:
        """Tovar nomidan brendni aniqlash."""
        if not title:
            return "Boshqa brendlar"

        title_upper = title.upper()
        for brand in KNOWN_BRANDS:
            pattern = rf"(?:\b|_|-){re.escape(brand.upper())}(?:\b|_|-)"
            if re.search(pattern, title_upper):
                return brand

        return "Boshqa brendlar"

    def clean_price(self, raw_price: Optional[str]) -> str:
        if not raw_price:
            return "Kelishilgan holda"
        text = raw_price.strip()
        lower = text.lower()
        if "по запросу" in lower or "kelishilgan" in lower or "договор" in lower:
            return "Kelishilgan holda"
        digits = re.sub(r"[^\d]", "", text)
        if digits:
            try:
                num = int(digits)
                return f"{num:,}".replace(",", " ") + " so'm"
            except ValueError:
                pass
        return text if text else "Kelishilgan holda"

    def parse_products_html(self, html: str, category_name: str) -> List[ScrapedProduct]:
        if "Ничего не найдено" in html or "hech narsa topilmadi" in html.lower():
            return []

        soup = BeautifulSoup(html, "html.parser")
        products: List[ScrapedProduct] = []

        cards = soup.select(".product-item, .proposal-item")
        if not cards:
            cards = soup.find_all("div", class_=lambda c: c and ("product-item" in c.split() or "product_card" in c))

        for card in cards:
            # Haqiqiy mahsulot sarlavhasi havolasini topish
            title_tag = None
            name = ""
            href = ""
            for a in card.find_all("a", href=True):
                a_href = a.get("href", "")
                a_text = a.get_text(strip=True)
                if (
                    a_text
                    and len(a_text) > 2
                    and "быстрый просмотр" not in a_text.lower()
                    and "/proposal/" in a_href
                    and "hot-view" not in a_href
                ):
                    title_tag = a
                    name = a_text
                    href = a_href
                    break

            if not title_tag or not name:
                continue

            product_url = urljoin(self.base_url, href)

            # Narx
            price_tag = card.select_one(".price, [class*='price']")
            raw_price = price_tag.get_text(strip=True) if price_tag else "Kelishilgan holda"
            price = self.clean_price(raw_price)

            # Rasm
            image_url = None
            imgs = card.find_all("img")
            for img in imgs:
                src = img.get("data-src") or img.get("data-original") or img.get("src")
                if src and not src.startswith("data:"):
                    image_url = urljoin(self.base_url, src)
                    break

            brand = self.extract_brand_name(name)

            products.append(
                ScrapedProduct(
                    name=name,
                    price=price,
                    category=category_name,
                    brand=brand,
                    image_url=image_url,
                    product_url=product_url,
                )
            )

        return products

    async def fetch_html(self, client: httpx.AsyncClient, url: str) -> Optional[str]:
        for attempt in range(3):
            try:
                resp = await client.get(url, headers=self.headers, timeout=25.0)
                if resp.status_code == 200:
                    return resp.text
                elif resp.status_code == 404:
                    return None
            except Exception:
                if attempt < 2:
                    await asyncio.sleep(1.0 + attempt)
        return None

    async def discover_all_proposal_categories(
        self, client: httpx.AsyncClient
    ) -> List[Tuple[str, str]]:
        """Barcha tovarli sahifalar (/proposals/...) ro'yxatini yig'ish."""
        logger.info("Katalog toifalari skanerlanmoqda...")
        main_html = await self.fetch_html(client, self.categories_url)
        if not main_html:
            return []

        soup = BeautifulSoup(main_html, "html.parser")
        proposal_urls: List[Tuple[str, str]] = []
        category_pages_to_check: List[Tuple[str, str]] = []
        seen = set()

        # 1. Asosiy /categories/ sahifasidagi havolalar
        for a in soup.find_all("a", href=True):
            href = a.get("href", "")
            text = a.get_text(strip=True)
            if not text or len(text) < 3 or "/uz/" in href:
                continue

            full_url = urljoin(self.base_url, href)
            if full_url in seen:
                continue

            if "/proposals/" in href:
                seen.add(full_url)
                proposal_urls.append((text, full_url))
            elif "/category/" in href and not href.endswith("/categories/"):
                seen.add(full_url)
                category_pages_to_check.append((text, full_url))

        logger.info(f"Topildi: {len(proposal_urls)} ta proposal toifasi, {len(category_pages_to_check)} ta umumiy kategoriya.")

        # 2. Umumiy kategoriyalar ichidagi ichki proposal toifalarini ochish
        semaphore = asyncio.Semaphore(self.max_concurrency)

        async def check_cat(cat_name: str, cat_url: str):
            async with semaphore:
                html = await self.fetch_html(client, cat_url)
            if not html:
                return []
            c_soup = BeautifulSoup(html, "html.parser")
            found = []
            for ca in c_soup.find_all("a", href=True):
                c_href = ca.get("href", "")
                c_text = ca.get_text(strip=True)
                if not c_text or len(c_text) < 3 or "/uz/" in c_href:
                    continue
                if "/proposals/" in c_href:
                    c_full = urljoin(self.base_url, c_href)
                    if c_full not in seen:
                        seen.add(c_full)
                        found.append((c_text, c_full))
            return found

        cat_tasks = [check_cat(name, url) for name, url in category_pages_to_check]
        sub_results = await asyncio.gather(*cat_tasks)
        for sub_list in sub_results:
            proposal_urls.extend(sub_list)

        logger.info(f"Jami yig'ilgan tovarli toifalar soni: {len(proposal_urls)}")
        return proposal_urls

    async def scrape_proposal_pages(
        self,
        client: httpx.AsyncClient,
        cat_name: str,
        cat_url: str,
        semaphore: asyncio.Semaphore,
        max_pages: int = 15,
    ) -> List[ScrapedProduct]:
        """Bitta proposal toifasini sahifalab skanerlash."""
        products: List[ScrapedProduct] = []
        page = 1

        while page <= max_pages:
            sep = "&" if "?" in cat_url else "?"
            page_url = f"{cat_url}{sep}page={page}" if page > 1 else cat_url

            async with semaphore:
                html = await self.fetch_html(client, page_url)

            if not html:
                break

            page_products = self.parse_products_html(html, category_name=cat_name)
            if not page_products:
                break

            current_urls = {p.product_url for p in products}
            new_unique = [p for p in page_products if p.product_url not in current_urls]
            if not new_unique and page > 1:
                break

            products.extend(new_unique)
            if len(page_products) < 5:
                break

            page += 1

        return products

    async def scrape_all(self, max_categories: Optional[int] = None) -> List[ScrapedProduct]:
        """Barcha mahsulotlarni asinxron skanerlash."""
        async with httpx.AsyncClient(verify=False, follow_redirects=True, timeout=20.0) as client:
            categories = await self.discover_all_proposal_categories(client)
            if max_categories:
                categories = categories[:max_categories]

            logger.info(f"{len(categories)} ta toifa bo'yicha tovarlarni yig'ish boshlandi...")
            semaphore = asyncio.Semaphore(self.max_concurrency)

            tasks = [
                self.scrape_proposal_pages(client, cat_name, cat_url, semaphore)
                for cat_name, cat_url in categories
            ]

            results = await asyncio.gather(*tasks)
            all_prods: List[ScrapedProduct] = []
            seen_urls = set()

            for prods in results:
                for p in prods:
                    if p.product_url not in seen_urls:
                        seen_urls.add(p.product_url)
                        all_prods.append(p)

            logger.info(f"Jami yig'ilgan noyob mahsulotlar soni: {len(all_prods)}")
            return all_prods

    def export_to_excel(
        self, products: List[ScrapedProduct], filename: str = "products_catalog.xlsx"
    ):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Mahsulotlar"

        headers = [
            "№",
            "Tovar Nomi",
            "Brend",
            "Toifa",
            "Narxi",
            "Rasm URL",
            "Tovar Havolasi",
        ]
        ws.append(headers)

        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

        for col_idx, col in enumerate(ws.iter_cols(min_row=1, max_row=1, min_col=1, max_col=len(headers)), start=1):
            cell = col[0]
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for idx, p in enumerate(products, start=1):
            ws.append([
                idx,
                p.name,
                p.brand,
                p.category,
                p.price,
                p.image_url or "",
                p.product_url,
            ])

        column_widths = {
            "A": 8,
            "B": 45,
            "C": 18,
            "D": 30,
            "E": 20,
            "F": 45,
            "G": 45,
        }
        for col_letter, width in column_widths.items():
            ws.column_dimensions[col_letter].width = width

        wb.save(filename)
        logger.info(f"Excel fayl saqlandi: {filename}")

    async def save_to_db(self, products: List[ScrapedProduct], db: DatabaseManager) -> int:
        brand_cache = {}
        category_cache = {}
        saved_count = 0

        for p in products:
            b_name = p.brand or "Boshqa brendlar"
            if b_name not in brand_cache:
                slug = re.sub(r"[^\w-]", "", b_name.lower().replace(" ", "-")) or "brand"
                brand_obj = await db.add_brand(name=b_name, slug=slug)
                brand_cache[b_name] = brand_obj.id
            brand_id = brand_cache[b_name]

            c_name = p.category or "Umumiy"
            cache_key = (c_name, brand_id)
            if cache_key not in category_cache:
                cat_obj = await db.add_category(name=c_name, brand_id=brand_id)
                category_cache[cache_key] = cat_obj.id
            category_id = category_cache[cache_key]

            await db.add_product(
                name=p.name,
                price=p.price,
                brand_id=brand_id,
                category_id=category_id,
                image_url=p.image_url,
                description=f"{p.name} - Control Pribor do'koni mahsuloti.",
                product_url=p.product_url,
            )
            saved_count += 1

        logger.info(f"Bazaga muvaffaqiyatli saqlangan tovarlar soni: {saved_count}")
        return saved_count


if __name__ == "__main__":
    from controlpribor_bot.config import Config

    async def main():
        config = Config.load()
        db = DatabaseManager(f"sqlite+aiosqlite:///{config.db_path}")
        await db.init_db()
        scraper = ControlPriborScraper()
        excel_path = os.path.join(os.path.dirname(__file__), "products_catalog.xlsx")
        logger.info("Control Pribor saytidan mahsulotlarni yig'ish boshlandi...")
        products = await scraper.scrape_all()
        scraper.export_to_excel(products, excel_path)
        await scraper.save_to_db(products, db)
        await db.close()

    asyncio.run(main())
